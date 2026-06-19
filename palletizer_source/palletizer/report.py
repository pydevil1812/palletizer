"""
Tabular output and exporters (Excel, PDF, printable HTML).
"""
from __future__ import annotations

import base64
import os
from datetime import datetime
from typing import Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate,
                                 Spacer, Table, TableStyle)

from .models import StackingResult


# ---------------------------------------------------------------------------
# Tabular data
# ---------------------------------------------------------------------------

def boxes_dataframe(result: StackingResult) -> pd.DataFrame:
    rows = [{
        "Box #": b.box_id,
        "Layer": b.layer + 1,
        "X (mm)": round(b.x, 1),
        "Y (mm)": round(b.y, 1),
        "Z (mm)": round(b.z, 1),
        "Dim X (mm)": round(b.dim_x, 1),
        "Dim Y (mm)": round(b.dim_y, 1),
        "Dim Z (mm)": round(b.dim_z, 1),
        "Orientation": b.orientation,
    } for b in result.placed_boxes]
    return pd.DataFrame(rows)


def layers_dataframe(result: StackingResult) -> pd.DataFrame:
    rows = [{
        "Layer": l.index + 1,
        "Z start (mm)": round(l.z_start, 1),
        "Layer height (mm)": round(l.dim_z, 1),
        "Boxes in layer": l.count,
        "Footprint per box (mm)": f"{l.dim_x:.0f} x {l.dim_y:.0f}",
        "Layer weight (kg)": round(l.weight, 1),
        "Orientation": l.orientation,
    } for l in result.layers]
    return pd.DataFrame(rows)


def summary_dict(result: StackingResult) -> Dict[str, str]:
    cfg = result.config
    return {
        "Pallet": cfg.pallet.name,
        "Box": cfg.box.name,
        "Total boxes": str(result.total_boxes),
        "Number of layers": str(len(result.layers)),
        "Boxes per layer (typical)": str(result.layers[0].count) if result.layers else "0",
        "Total gross weight (kg)": f"{result.total_weight:.1f}",
        "Pallet load capacity (kg)": f"{cfg.pallet.load_capacity:.1f}",
        "Weight utilization (%)": f"{result.weight_utilization_pct:.1f}",
        "Total height incl. pallet (mm)": f"{result.total_height:.1f}",
        "Max allowed stack height (mm)": f"{cfg.max_stack_height:.1f}",
        "Height utilization (%)": f"{result.height_utilization_pct:.1f}",
        "Avg. footprint fill per layer (%)": f"{result.footprint_fill_pct:.1f}",
        "Volumetric fill factor (%)": f"{result.volume_fill_pct:.1f}",
        "Report generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1):
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
    for c_idx, col in enumerate(df.columns, start=1):
        width = max(12, min(28, int(df[col].astype(str).map(len).max() if len(df) else 10) + 4))
        ws.column_dimensions[get_column_letter(c_idx)].width = width


def export_excel(result: StackingResult, path: str) -> str:
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum["A1"] = "Pallet Loading — Summary"
    ws_sum["A1"].font = TITLE_FONT
    for i, (k, v) in enumerate(summary_dict(result).items(), start=3):
        ws_sum.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws_sum.cell(row=i, column=2, value=v)
    ws_sum.column_dimensions["A"].width = 34
    ws_sum.column_dimensions["B"].width = 24

    ws_rec = wb.create_sheet("Recommendations")
    ws_rec["A1"] = "Recommendations"
    ws_rec["A1"].font = TITLE_FONT
    if result.recommendations:
        for i, rec in enumerate(result.recommendations, start=3):
            ws_rec.cell(row=i, column=1, value=f"{i - 2}. {rec}")
    else:
        ws_rec.cell(row=3, column=1, value="No recommendations — configuration looks efficient.")
    ws_rec.column_dimensions["A"].width = 110

    ws_layers = wb.create_sheet("Layers")
    _write_df(ws_layers, layers_dataframe(result))

    ws_boxes = wb.create_sheet("Box Coordinates")
    _write_df(ws_boxes, boxes_dataframe(result))

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------

def export_pdf(result: StackingResult, images: Dict[str, str], path: str) -> str:
    doc = SimpleDocTemplate(path, pagesize=A4,
                             topMargin=18 * mm, bottomMargin=16 * mm,
                             leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=20)
    h2 = styles["Heading2"]
    body = styles["BodyText"]
    story = []

    story.append(Paragraph("Pallet Loading Report", title_style))
    story.append(Paragraph(f"{result.config.pallet.name} / {result.config.box.name}", body))
    story.append(Spacer(1, 10))

    summary = summary_dict(result)
    table_data = [[k, v] for k, v in summary.items()]
    t = Table(table_data, colWidths=[80 * mm, 80 * mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F8")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    story.append(Paragraph("Recommendations", h2))
    if result.recommendations:
        for rec in result.recommendations:
            story.append(Paragraph(f"• {rec}", body))
            story.append(Spacer(1, 4))
    else:
        story.append(Paragraph("No recommendations — configuration looks efficient.", body))

    story.append(PageBreak())
    story.append(Paragraph("Diagrams", h2))
    for key, caption in (("top_view", "Top view by layer pattern"),
                          ("side_view", "Front / side elevation"),
                          ("view_3d", "3D view")):
        img_path = images.get(key)
        if img_path and os.path.exists(img_path):
            story.append(Paragraph(caption, styles["Heading3"]))
            story.append(Image(img_path, width=170 * mm, height=170 * mm * 0.72, kind="proportional"))
            story.append(Spacer(1, 10))

    story.append(PageBreak())
    story.append(Paragraph("Layer breakdown", h2))
    ldf = layers_dataframe(result)
    if not ldf.empty:
        data = [list(ldf.columns)] + ldf.astype(str).values.tolist()
        lt = Table(data, repeatRows=1)
        lt.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FB")]),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]))
        story.append(lt)

    story.append(PageBreak())
    story.append(Paragraph("Box coordinates", h2))
    bdf = boxes_dataframe(result)
    if not bdf.empty:
        data = [list(bdf.columns)] + bdf.astype(str).values.tolist()
        bt = Table(data, repeatRows=1)
        bt.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 6.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FB")]),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]))
        story.append(bt)
    else:
        story.append(Paragraph("No boxes were placed.", body))

    doc.build(story)
    return path


# ---------------------------------------------------------------------------
# Printable HTML export (self-contained, images inlined as base64)
# ---------------------------------------------------------------------------

def _img_to_base64(path: str) -> str:
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode("ascii")


def export_html(result: StackingResult, images: Dict[str, str], path: str) -> str:
    summary = summary_dict(result)
    rows_html = "".join(f"<tr><th>{k}</th><td>{v}</td></tr>" for k, v in summary.items())

    rec_html = "".join(f"<li>{r}</li>" for r in result.recommendations) or \
        "<li>No recommendations — configuration looks efficient.</li>"

    def img_tag(key):
        p = images.get(key)
        if p and os.path.exists(p):
            return f'<img src="data:image/png;base64,{_img_to_base64(p)}" class="diagram"/>'
        return ""

    ldf = layers_dataframe(result)
    bdf = boxes_dataframe(result)

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pallet Loading Report</title>
<style>
  body {{ font-family: Arial, Helvetica, sans-serif; color: #222; margin: 24px; }}
  h1 {{ font-size: 22px; margin-bottom: 0; }}
  h2 {{ font-size: 16px; border-bottom: 2px solid #1F4E78; color: #1F4E78; padding-bottom: 4px; margin-top: 28px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 12px; margin-top: 8px; }}
  th, td {{ border: 1px solid #ccc; padding: 5px 8px; text-align: left; }}
  table.summary th {{ background: #EEF3F8; width: 280px; }}
  table.data th {{ background: #1F4E78; color: white; }}
  table.data tr:nth-child(even) {{ background: #F5F8FB; }}
  .diagram {{ max-width: 100%; margin: 10px 0; border: 1px solid #ddd; }}
  ul {{ font-size: 13px; }}
  .meta {{ color: #666; font-size: 12px; }}
  @media print {{
    body {{ margin: 8mm; }}
    h2 {{ page-break-before: auto; }}
    table.data {{ page-break-inside: auto; }}
    tr {{ page-break-inside: avoid; }}
  }}
</style>
</head>
<body>
  <h1>Pallet Loading Report</h1>
  <div class="meta">{result.config.pallet.name} / {result.config.box.name}</div>

  <h2>Summary</h2>
  <table class="summary">{rows_html}</table>

  <h2>Recommendations</h2>
  <ul>{rec_html}</ul>

  <h2>Diagrams</h2>
  {img_tag('top_view')}
  {img_tag('side_view')}
  {img_tag('view_3d')}

  <h2>Layer breakdown</h2>
  {ldf.to_html(index=False, classes="data", border=0) if not ldf.empty else "<p>No layers.</p>"}

  <h2>Box coordinates</h2>
  {bdf.to_html(index=False, classes="data", border=0) if not bdf.empty else "<p>No boxes placed.</p>"}

</body>
</html>"""

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path
